"""Document extraction for directory-based datasets.

A dataset is a directory of files (PDF, Excel, CSV, XML, text …). These helpers
scan the directory and extract plain text so it can be chunked, embedded and
indexed into the RAG knowledge base.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.core.logging import get_logger

logger = get_logger("discoveryx.documents")

TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log"}
DATA_EXTS = {".json", ".jsonl", ".csv", ".tsv", ".xml", ".html", ".htm"}
SHEET_EXTS = {".xlsx", ".xlsm", ".xls"}
DOC_EXTS = {".pdf", ".docx"}
SUPPORTED_EXTS = TEXT_EXTS | DATA_EXTS | SHEET_EXTS | DOC_EXTS
# Platform metadata files are not dataset content.
METADATA_NAMES = {"dataset.json", ".index_state.json"}


def _is_content(p: Path) -> bool:
    return p.is_file() and p.name not in METADATA_NAMES


def extract_text(path: Path, *, max_chars: int = 200_000) -> str:
    """Extract plain text from a file. Returns ``""`` if unsupported/unreadable."""
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="replace")[:max_chars]
        if ext == ".json":
            data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
            return json.dumps(data, ensure_ascii=False, indent=2)[:max_chars]
        if ext == ".jsonl":
            lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
            return "\n".join(lines[:2000])[:max_chars]
        if ext in {".csv", ".tsv"}:
            delim = "\t" if ext == ".tsv" else ","
            out: list[str] = []
            with open(path, encoding="utf-8", errors="replace", newline="") as fh:
                reader = csv.reader(fh, delimiter=delim)
                for i, row in enumerate(reader):
                    if i > 5000:
                        break
                    out.append(" | ".join(row))
            return "\n".join(out)[:max_chars]
        if ext in {".xml", ".html", ".htm"}:
            raw = path.read_text(encoding="utf-8", errors="replace")
            try:
                from lxml import etree

                parser = etree.XMLParser(recover=True, huge_tree=True)
                root = etree.fromstring(raw.encode("utf-8", "replace"), parser=parser)
                text = " ".join("".join(root.itertext()).split())
                return text[:max_chars]
            except Exception:  # noqa: BLE001
                return raw[:max_chars]
        if ext in SHEET_EXTS:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"# sheet: {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 3000:
                        break
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" | ".join(cells))
            wb.close()
            return "\n".join(out)[:max_chars]
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = []
            for page in reader.pages[:200]:
                pages.append(page.extract_text() or "")
            return "\n\n".join(pages)[:max_chars]
        if ext == ".docx":
            from docx import Document

            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)[:max_chars]
    except Exception as exc:  # noqa: BLE001 - never fail a whole scan on one file
        logger.warning("failed to extract %s: %s", path, exc)
    return ""


def scan_directory(root: Path) -> dict[str, Any]:
    """Summarise a dataset directory: file count, size and extension histogram."""
    types: Counter[str] = Counter()
    n_files = 0
    size = 0
    for p in root.rglob("*"):
        if _is_content(p):
            n_files += 1
            size += p.stat().st_size
            types[p.suffix.lower() or "(none)"] += 1
    return {"n_files": n_files, "size_bytes": size, "file_types": dict(types)}


def iter_documents(root: Path, *, max_files: int = 2000) -> Iterator[dict[str, Any]]:
    """Yield ``{id, text, metadata}`` for every readable file under ``root``."""
    count = 0
    for p in sorted(root.rglob("*")):
        if not p.is_file() or p.suffix.lower() not in SUPPORTED_EXTS:
            continue
        text = extract_text(p)
        if not text.strip():
            continue
        rel = p.relative_to(root).as_posix()
        yield {
            "id": f"{root.name}:{rel}",
            "text": text,
            "metadata": {
                "file": rel,
                "ext": p.suffix.lower(),
                "dataset": root.name,
                "size_bytes": p.stat().st_size,
            },
        }
        count += 1
        if count >= max_files:
            break


def list_files(root: Path, *, limit: int = 500) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for p in sorted(root.rglob("*")):
        if not _is_content(p):
            continue
        st = p.stat()
        items.append(
            {
                "name": p.name,
                "rel_path": p.relative_to(root).as_posix(),
                "ext": p.suffix.lower() or "(none)",
                "size_bytes": st.st_size,
                "modified_at": datetime.fromtimestamp(st.st_mtime, UTC).isoformat(timespec="seconds"),
            }
        )
        if len(items) >= limit:
            break
    return items
